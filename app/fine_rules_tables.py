"""Owner mapping workbook to fine-rules tables: VAT codes, location master, supplier sites.

The workbook is private. It is read in place and turned into the ``vat_codes``,
``location_master`` and ``supplier_sites`` fields of the fine-rules config. The
report carries counts, sheet names and row numbers only, never a cell value.
Rows that cannot be used are flagged with their row number and reason, never
dropped silently.

    python -m app.fine_rules_tables WORKBOOK --out PRIVATE.json [--merge CURRENT.json]
"""

import argparse
import json
import os
import re
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from app.fine_rules import STORE, WAREHOUSE, to_decimal
from app.mapping_headers import SHEETS

VAT_SHEET, LOCATION_SHEET, SITE_SHEET = "VAT CODES", "LOCATIONS", "SUPPLIER SITES"
LOCATION_TYPES = {"STORE": STORE, "W/H": WAREHOUSE}
SITE_STATUSES = {"Active", "Inactive"}
MAX_ROWS = 25_000
MAX_COLUMNS = 64


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _rows(workbook, name):
    """Header-keyed rows with their sheet row number; fully blank rows are skipped."""
    if isinstance(workbook, dict):
        yield from workbook[name]
        return
    sheet = next((ws for ws in workbook.worksheets if ws.title.strip().upper() == name), None)
    if sheet is None:
        raise ValueError(f"Sheet {name} is missing")
    rows = sheet.iter_rows(values_only=True)
    header = [_cell(x).upper() for x in next(rows, ())]
    for number, values in enumerate(rows, 2):
        if all(_cell(v) == "" for v in values):
            continue
        yield number, {k: v for k, v in zip(header, values) if k}


def _need(row, *columns):
    missing = [c for c in columns if c not in row]
    if missing:
        raise ValueError(f"Columns missing: {', '.join(missing)}")


def vat_codes(workbook, flag):
    """Owner instruction: only VAT type C whose code contains PV. Region = VAT_REGION_NAME without 'Vat Region'."""
    kept, excluded = {}, 0
    for number, row in _rows(workbook, VAT_SHEET):
        _need(row, "VAT_TYPE", "VAT_CODE", "VAT_REGION_NAME", "VAT_RATE", "ACTIVE_DATE")
        kind, code = _cell(row["VAT_TYPE"]).upper(), _cell(row["VAT_CODE"]).upper()
        if not (kind == "C" and "PV" in code):
            excluded += 1
            continue
        region = re.sub(r"\s*VAT\s+REGION\s*$", "", _cell(row["VAT_REGION_NAME"]).upper()).strip()
        rate = to_decimal(row["VAT_RATE"])
        active = row["ACTIVE_DATE"]
        active = active.date() if isinstance(active, datetime) else active
        if not region or rate is None or rate < 0 or not isinstance(active, date):
            flag(VAT_SHEET, number, "region, rate or active date missing or invalid")
        elif region in kept:
            flag(VAT_SHEET, number, "second C/PV code for the same VAT region")
        else:
            kept[region] = {"code": code, "region": region, "rate": str(rate), "active_from": active.isoformat()}
    return list(kept.values()), excluded


def locations(workbook, flag):
    """Location master: type from WH/STORE, market and country from COUNTRY (R-011, R-027, ALG-014)."""
    master = {}
    for number, row in _rows(workbook, LOCATION_SHEET):
        _need(row, "LOCATION", "WH/STORE", "COUNTRY")
        location, country = _cell(row["LOCATION"]), _cell(row["COUNTRY"]).upper()
        kind = LOCATION_TYPES.get(_cell(row["WH/STORE"]).upper())
        entity = _cell(row.get("ENTITY AND CURENCY", row.get("ENTITY AND CURRENCY"))).upper()
        if not location or not kind or not country:
            flag(LOCATION_SHEET, number, "location id, type (STORE or W/H) or country missing")
        elif location in master:
            flag(LOCATION_SHEET, number, "location id repeats")
        else:
            master[location] = {"type": kind, "market": country, "country": country, "entity_currency": entity}
    return master


def supplier_sites(workbook, flag):
    """Supplier site id, its one currency and status, with supplier code, name and site name for the SUP-001
    bridge. Rows without a site id cannot be keyed and are counted."""
    sites, without_id = {}, 0
    for number, row in _rows(workbook, SITE_SHEET):
        _need(row, "SUPPLIER SITE", "STATUS DESCRIPTION", "CURRENCY")
        site = _cell(row["SUPPLIER SITE"])
        if not site:
            without_id += 1
            flag(SITE_SHEET, number, "supplier site id missing")
            continue
        currency, status = _cell(row["CURRENCY"]).upper(), _cell(row["STATUS DESCRIPTION"]).title()
        if not re.fullmatch(r"[A-Z]{3}", currency) or status not in SITE_STATUSES:
            flag(SITE_SHEET, number, "currency is not a 3-letter code or status is not Active/Inactive")
        elif site in sites:
            flag(SITE_SHEET, number, "supplier site id repeats")
        else:
            # SUP-001 bridge fields; the config stays private (names never leave it).
            sites[site] = {"supplier_site": site, "currency": currency, "status": status,
                           "supplier_code": _cell(row.get("SUPPLIER CODE")), "supplier_name": _cell(row.get("SUPPLIER NAME")),
                           "site_name": _cell(row.get("SUPPLIER SITE NAME"))}
    return list(sites.values()), without_id


def load(path):
    import openpyxl

    # Do not use cached formula results: mapping values must be literal and reviewable.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    flagged = []

    def flag(sheet, row, reason):
        flagged.append({"sheet": sheet, "row": row, "reason": reason})

    try:
        rows, sheets = {}, {}
        for name, spec in SHEETS.items():
            rows[name] = []
            counts = sheets[name] = {"rows": 0, "accepted": 0, "excluded": 0, "flagged": 0}
            matches = [ws for ws in workbook.worksheets if ws.title.strip().upper() == name]
            if len(matches) != 1:
                flag(name, 1, "required sheet missing or duplicated")
                continue
            sheet = matches[0]
            # Reset untrusted dimensions: iterating otherwise follows forged Excel max-row metadata.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            header = next(iterator, ())
            aliases = {alias.strip().upper(): field for field, names in spec["headers"].items() for alias in names}
            fields = [aliases.get(_cell(cell.value).upper()) for cell in header]
            missing = [field for field in spec["required"] if field not in fields]
            duplicate = len([f for f in fields if f]) != len({f for f in fields if f})
            if len(header) > MAX_COLUMNS:
                raise ValueError("Mapping sheets support at most 64 columns")
            if missing:
                flag(name, 1, "required columns missing: " + ", ".join(missing))
            if duplicate:
                flag(name, 1, "multiple columns map to the same field")
            for number, cells in enumerate(iterator, 2):
                if number > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                    raise ValueError("Mapping sheets support at most 25000 rows and 64 columns")
                if all(_cell(cell.value) == "" for cell in cells):
                    continue
                counts["rows"] += 1
                if any(cell.data_type in ("f", "e") for cell in cells):
                    flag(name, number, "formula or Excel error; replace with literal values")
                elif not missing and not duplicate:
                    values = [cell.value for cell in cells]
                    rows[name].append((number, {field: values[i] if i < len(values) else None
                                               for i, field in enumerate(fields) if field}))
        codes, excluded = vat_codes(rows, flag)
        master = locations(rows, flag)
        sites, without_id = supplier_sites(rows, flag)
        for name, accepted in ((VAT_SHEET, len(codes)), (LOCATION_SHEET, len(master)), (SITE_SHEET, len(sites))):
            sheets[name]["accepted"] = accepted
            sheets[name]["flagged"] = len({f["row"] for f in flagged if f["sheet"] == name and f["row"] > 1})
            if not accepted and not any(f["sheet"] == name for f in flagged):
                flag(name, 1, "sheet has no usable mapping rows")
        sheets[VAT_SHEET]["excluded"] = excluded
    finally:
        workbook.close()
    tables = {"vat_codes": codes, "location_master": master, "supplier_sites": sites}
    report = {
        "vat_codes_c_pv": len(codes), "vat_rows_excluded": excluded, "locations": len(master),
        "locations_by_type": {k: sum(1 for r in master.values() if r["type"] == k) for k in (STORE, WAREHOUSE)},
        "supplier_sites": len(sites), "supplier_sites_active": sum(1 for s in sites if s["status"] == "Active"),
        "supplier_sites_inactive": sum(1 for s in sites if s["status"] == "Inactive"),
        "supplier_rows_without_site_id": without_id,
        "site_currencies": len({s["currency"] for s in sites}), "flagged": flagged,
        "sheets": sheets, "valid": not flagged,
    }
    return tables, report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--out", type=Path, required=True, help="private JSON config to write (mode 0600)")
    parser.add_argument("--merge", type=Path, help="current fine-rules config JSON whose other fields are kept")
    parser.add_argument("--version", default=None, help="config version label")
    args = parser.parse_args(argv)
    tables, report = load(args.workbook)
    if not report["valid"]:
        json.dump(report, sys.stdout, indent=1)
        print()
        return 1
    config = json.loads(args.merge.read_text()) if args.merge else {}
    config.update(tables)
    if args.version:
        config["version"] = args.version
    from app.fine_rules import RulesConfig

    RulesConfig.from_dict(config)
    descriptor = os.open(args.out, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w") as handle:
        json.dump(config, handle, default=lambda v: str(v) if isinstance(v, Decimal) else v)
    json.dump(report, sys.stdout, indent=1)
    print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
